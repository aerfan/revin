"""Recalculate every cache-free fuzz workbook in LibreOffice and compare with the engine."""
import json, os, shutil, subprocess, sys
from openpyxl import load_workbook

OUT = os.path.join(os.path.dirname(__file__), 'out', 'fuzz')
RECALC = '/mnt/skills/public/xlsx/scripts'
manifest = json.load(open(os.path.join(OUT, 'manifest.json')))
targets = [m for m in manifest if m['noCache']]
problems, compared = [], 0
for m in targets:
    work = m['f'].replace('.xlsx', '-recalc.xlsx')
    shutil.copy(m['f'], work)
    r = subprocess.run(['python3', 'recalc.py', work, '90'], cwd=RECALC, capture_output=True, text=True, timeout=200)
    try:
        info = json.loads(r.stdout[r.stdout.index('{'):])
    except Exception:
        problems.append([m['tag'], 'recalc failed: ' + (r.stdout + r.stderr)[-200:]]); continue
    if info.get('total_errors'):
        problems.append([m['tag'], 'formula errors: ' + json.dumps(info.get('error_summary'))[:200]])
    wb = load_workbook(work, data_only=True)
    bs, ss = wb['Breakdown'], wb['Summary']
    total = next(bs.cell(r, 9).value for r in range(2, bs.max_row + 1) if bs.cell(r, 3).value == 'Total hours')
    if total is None or abs(total - m['total']) > 1e-6 * max(1, m['total']):
        problems.append([m['tag'], f'recalculated total {total} != engine {m["total"]}'])
    summary_total = next((ss.cell(r, 2).value for r in range(1, ss.max_row + 1) if ss.cell(r, 1).value == 'Total hours'), None)
    if summary_total is None or abs(summary_total - m['total']) > 1e-6 * max(1, m['total']):
        problems.append([m['tag'], f'summary total {summary_total} != engine {m["total"]}'])
    for name, ptotal in m['phases']:
        row = next((r for r in range(1, ss.max_row + 1) if ss.cell(r, 1).value == name), None)
        if row is None:
            continue
        v = ss.cell(row, 4).value
        if v is None or abs(v - ptotal) > 1e-6 * max(1, ptotal):
            problems.append([m['tag'], f'phase {name}: {v} != {ptotal}'])
    compared += 1
print(json.dumps({'workbooks': len(targets), 'compared': compared, 'problems': problems}, indent=1))
sys.exit(1 if problems else 0)
